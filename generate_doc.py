import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_training_document(filename):
    doc = Document()

    # Configurar márgenes (1 pulgada = 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # Colores temáticos
    COLOR_PRIMARY = RGBColor(19, 102, 170)      # Azul petróleo / tech
    COLOR_SECONDARY = RGBColor(19, 142, 102)    # Verde MongoDB
    COLOR_DARK = RGBColor(40, 50, 60)           # Texto oscuro
    COLOR_MUTED = RGBColor(90, 105, 120)        # Texto secundario
    COLOR_CODE_BG = "F4F6F8"                    # Fondo código
    COLOR_CALLOUT_BG = "EBF5FB"                 # Fondo llamada

    # Estilo Normal
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_DARK
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

    # Funciones de utilidad para estilos XML
    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def add_title(text, subtitle=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(22)
        run.bold = True
        run.font.color.rgb = COLOR_PRIMARY

        if subtitle:
            p_sub = doc.add_paragraph()
            p_sub.paragraph_format.space_after = Pt(14)
            run_sub = p_sub.add_run(subtitle)
            run_sub.font.name = 'Segoe UI'
            run_sub.font.size = Pt(11)
            run_sub.font.color.rgb = COLOR_MUTED

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(14)
        run.bold = True
        run.font.color.rgb = COLOR_PRIMARY

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI Semibold'
        run.font.size = Pt(12)
        run.bold = True
        run.font.color.rgb = COLOR_SECONDARY

    def add_callout(text, label="NOTA CLAVE:"):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(6.8)
        
        cell = table.cell(0, 0)
        set_cell_background(cell, COLOR_CALLOUT_BG)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        
        # Borde izquierdo grueso
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="1366AA"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.line_spacing = 1.15
        
        r_label = cp.add_run(f"{label} ")
        r_label.bold = True
        r_label.font.name = 'Segoe UI'
        r_label.font.size = Pt(9.5)
        r_label.font.color.rgb = COLOR_PRIMARY
        
        r_text = cp.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = COLOR_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def add_code_block(code_str):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(6.8)
        
        cell = table.cell(0, 0)
        set_cell_background(cell, COLOR_CODE_BG)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        
        # Bordes sutiles
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="D5DBDB"/>
                <w:left w:val="single" w:sz="18" w:space="0" w:color="138E66"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="D5DBDB"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="D5DBDB"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)

        lines = code_str.strip().split("\n")
        cell.paragraphs[0].text = "" # limpiar primer parrafo
        for i, line in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(30, 40, 50)

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # --- CONTENIDO DEL DOCUMENTO ---

    # Portada / Cabecera compacta
    add_title("Guía Práctica de Capacitación: MongoDB y Mongoose", 
              "Kit de iniciación rápida para nuevos integrantes del equipo de desarrollo")

    add_callout(
        "Documentación oficial de referencia sobre documentos MongoDB: "
        "https://www.mongodb.com/es/docs/manual/core/document/ - "
        "Consúltala para profundizar en la estructura BSON y especificaciones oficiales.",
        label="DOCUMENTACIÓN OFICIAL:"
    )

    # 1. FUNDAMENTOS
    add_h1("1. Fundamentos de Bases de Datos NoSQL y MongoDB")
    p = doc.add_paragraph(
        "A diferencia de las bases de datos relacionales (SQL) tradicionales basadas en tablas rígidas y columnas fijas, "
        "MongoDB es una base de datos NoSQL orientada a documentos. Almacena la información en formato BSON "
        "(Binary JSON), lo que permite esquemas flexibles, campos anidados y arreglos dentro del mismo registro."
    )

    # Tabla de equivalencias SQL vs MongoDB
    add_h2("Equivalencia de conceptos: SQL vs. MongoDB")
    table_eq = doc.add_table(rows=5, cols=3)
    table_eq.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_eq.columns[0].width = Inches(1.8)
    table_eq.columns[1].width = Inches(2.2)
    table_eq.columns[2].width = Inches(2.8)

    headers = ["Concepto SQL (Relacional)", "Concepto MongoDB", "Descripción"]
    for j, h in enumerate(headers):
        cell = table_eq.cell(0, j)
        set_cell_background(cell, "1366AA")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        run.bold = True
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("Base de datos (Database)", "Base de datos (Database)", "Contenedor lógico de la información."),
        ("Tabla (Table)", "Colección (Collection)", "Agrupación de documentos (no exige esquema idéntico)."),
        ("Fila / Registro (Row)", "Documento (Document)", "Objeto individual en formato BSON/JSON con un campo _id único."),
        ("Columna (Column)", "Campo (Field)", "Par clave-valor dentro de un documento (admite arrays y objetos).")
    ]

    for i, row in enumerate(data):
        row_cells = table_eq.rows[i+1].cells
        bg_col = "F9FAFB" if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate(row):
            cell = row_cells[j]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(val)
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 2. INSTALACIÓN
    add_h1("2. Instalación y Puesta en Marcha")
    doc.add_paragraph(
        "Para trabajar con MongoDB y Node.js en nuestro entorno, disponemos de dos caminos principales:"
    )

    p_inst1 = doc.add_paragraph()
    p_inst1.paragraph_format.left_indent = Inches(0.2)
    p_inst1.paragraph_format.space_after = Pt(2)
    r = p_inst1.add_run("• Opción A - MongoDB Atlas (Recomendada para iniciar): ")
    r.bold = True
    p_inst1.add_run("Base de datos en la nube gratuita (Free Shared Cluster). No requiere instalar software en tu máquina; solo creas una cuenta, obtienes la cadena de conexión (URI) y listo.")

    p_inst2 = doc.add_paragraph()
    p_inst2.paragraph_format.left_indent = Inches(0.2)
    p_inst2.paragraph_format.space_after = Pt(4)
    r2 = p_inst2.add_run("• Opción B - Local con MongoDB Community Server + Compass: ")
    r2.bold = True
    p_inst2.add_run("Descargas el instalador de MongoDB Community Edition y MongoDB Compass (interfaz gráfica oficial para explorar y consultar datos visualmente).")

    add_h2("Instalación de Mongoose y Conexión en Node.js")
    doc.add_paragraph(
        "Mongoose es el ODM (Object Data Modeling) estándar de la industria para Node.js. "
        "Nos permite modelar datos, agregar validaciones y gestionar relaciones de manera declarativa y segura."
    )

    add_code_block("""// 1. En la terminal de tu proyecto:
npm install mongoose

// 2. Archivo db.js (Conexión básica):
const mongoose = require('mongoose');

const connectDB = async () => {
  try {
    await mongoose.connect('mongodb://localhost:27017/mi_empresa');
    console.log('MongoDB conectado exitosamente');
  } catch (error) {
    console.error('Error al conectar a MongoDB:', error.message);
    process.exit(1);
  }
};

module.exports = connectDB;""")

    # 3. VALIDACIÓN DE DATOS
    add_h1("3. Validación de Datos: ¿Soporte Nativo o Librerías?")
    doc.add_paragraph(
        "Una duda recurrente al ingresar es: ¿MongoDB valida datos por sí mismo o requiere librerías externas?"
    )

    add_callout(
        "RESPUESTA DIRECTA: Ambos mecanismos existen y se complementan. MongoDB cuenta con validación "
        "NATIVA a nivel de motor de base de datos, pero en el stack de desarrollo con Node.js utilizamos "
        "Mongoose a nivel de APLICACIÓN por conveniencia, tipado estricto y captura temprana de errores.",
        label="CLARIFICACIÓN:"
    )

    add_h2("A. Validación Nativa de MongoDB (JSON Schema a nivel BD)")
    doc.add_paragraph(
        "MongoDB soporta validación nativa directamente mediante el operador $jsonSchema al crear o modificar colecciones. "
        "La base de datos rechaza cualquier inserción que no cumpla las reglas, sin importar qué cliente o lenguaje envíe la petición:"
    )

    add_code_block("""// Ejemplo de validación nativa ejecutada en MongoDB Shell / Compass:
db.createCollection("usuarios", {
  validator: {
    $jsonSchema: {
      bsonType: "object",
      required: ["nombre", "email", "edad"],
      properties: {
        nombre: { bsonType: "string", description: "Debe ser texto y es obligatorio" },
        email:  { bsonType: "string", pattern: "^.+@.+$", description: "Email válido" },
        edad:   { bsonType: "int", minimum: 18, description: "Debe ser mayor o igual a 18" }
      }
    }
  }
});""")

    add_h2("B. Validación a Nivel Aplicación con Mongoose (Lo que usamos en el código)")
    doc.add_paragraph(
        "En la práctica diaria del equipo, definimos las reglas en el Schema de Mongoose. "
        "Esto evita viajes innecesarios a la base de datos si el payload es inválido y genera mensajes de error claros para la API:"
    )

    add_code_block(r"""const mongoose = require('mongoose');

const usuarioSchema = new mongoose.Schema({
  nombre: {
    type: String,
    required: [true, 'El nombre es obligatorio'],
    trim: true,
    minlength: [2, 'Mínimo 2 caracteres']
  },
  email: {
    type: String,
    required: true,
    unique: true,
    lowercase: true,
    match: [/^\S+@\S+\.\S+$/, 'Formato de correo no válido']
  },
  rol: {
    type: String,
    enum: ['ADMIN', 'DEVELOPER', 'GUEST'],
    default: 'DEVELOPER'
  },
  edad: {
    type: Number,
    min: [18, 'Debe ser mayor de edad']
  }
}, { timestamps: true });

module.exports = mongoose.model('Usuario', usuarioSchema);""")

    # 4. RELACIONES
    add_h1("4. Relaciones entre Documentos y Bases de Datos")
    doc.add_paragraph(
        "En MongoDB no existen las restricciones de clave foránea (FOREIGN KEY) automáticas ni el operador JOIN nativo de SQL. "
        "Existen dos estrategias principales para modelar relaciones: Embebido (Embedding) y Referenciado (Referencing)."
    )

    add_h2("Estrategia 1: Documentos Embebidos (Embedding / Desnormalización)")
    doc.add_paragraph(
        "Consiste en guardar el documento relacionado dentro del documento principal como un subdocumento o un arreglo de objetos. "
        "Se recomienda cuando los datos se consultan siempre juntos y la relación es 1 a 1 o 1 a pocos (ej. una orden con sus direcciones de entrega)."
    )

    add_code_block("""// Ejemplo de Embebido (Dirección dentro de Usuario):
const usuarioConDireccionSchema = new mongoose.Schema({
  nombre: String,
  direccion: {
    calle: String,
    ciudad: String,
    codigoPostal: String
  }
});
// Ventaja: 1 sola lectura a disco para traer usuario y su dirección. Excelente rendimiento.""")

    add_h2("Estrategia 2: Referenciado (Referencing / Normalización)")
    doc.add_paragraph(
        "Consiste en almacenar únicamente el _id del documento relacionado (tipo ObjectId) con la referencia (ref) al modelo destino. "
        "Se utiliza para relaciones 1 a Muchos grandes o Muchos a Muchos (N:M), evitando duplicación de datos y superación del límite de 16MB por documento."
    )

    add_code_block("""// 1. Modelo de Autor
const Autor = mongoose.model('Autor', new mongoose.Schema({ nombre: String }));

// 2. Modelo de Libro que referencia al Autor mediante ObjectId
const libroSchema = new mongoose.Schema({
  titulo: { type: String, required: true },
  autor: {
    type: mongoose.Schema.Types.ObjectId,
    ref: 'Autor', // Nombre del modelo al que apunta
    required: true
  }
});
const Libro = mongoose.model('Libro', libroSchema);

// 3. Consulta utilizando .populate() para unir los datos (similar a un JOIN):
const libroConAutor = await Libro.findOne({ titulo: 'Clean Code' }).populate('autor');
console.log(libroConAutor.autor.nombre); // Muestra los datos completos del Autor""")

    # Tabla resumen de cuándo usar cada uno
    add_h2("Resumen: ¿Cuándo embeber y cuándo referenciar?")
    table_rel = doc.add_table(rows=4, cols=3)
    table_rel.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_rel.columns[0].width = Inches(1.8)
    table_rel.columns[1].width = Inches(2.5)
    table_rel.columns[2].width = Inches(2.5)

    rel_headers = ["Criterio", "Embebido (Subdocumentos)", "Referenciado (ObjectId + Populate)"]
    for j, h in enumerate(rel_headers):
        cell = table_rel.cell(0, j)
        set_cell_background(cell, "138E66")
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        run.bold = True
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    rel_data = [
        ("Tipo de relación", "1 a 1, o 1 a pocos acotados (< 100).", "1 a muchos ilimitados, o Muchos a Muchos (N:M)."),
        ("Patrón de consulta", "Se leen y actualizan casi siempre juntos.", "Se consultan de forma independiente con frecuencia."),
        ("Rendimiento", "Máxima velocidad (1 sola consulta, sin joins).", "Requiere consultas adicionales o uso de $lookup / populate.")
    ]

    for i, row in enumerate(rel_data):
        row_cells = table_rel.rows[i+1].cells
        bg_col = "F9FAFB" if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate(row):
            cell = row_cells[j]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(val)
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 5. REGLAS DE ORO
    add_h1("5. Resumen de Reglas de Oro para el Equipo")
    bullet_points = [
        ("Diseña según las consultas: ", "En MongoDB modelamos los datos en función de cómo la aplicación va a consultar y mostrar la información, no solo de cómo se relacionan conceptualmente."),
        ("Límite de 16 MB por documento: ", "Un documento BSON nunca debe superar los 16 MB. Si una lista embebida puede crecer indefinidamente (ej. logs o comentarios infinitos), usa referencias obligatoriamente."),
        ("Valida siempre en Mongoose: ", "Define schemas completos con mensajes descriptivos para que los errores lleguen limpios a la capa de controladores y respuestas HTTP.")
    ]

    for title_b, desc_b in bullet_points:
        bp = doc.add_paragraph()
        bp.paragraph_format.left_indent = Inches(0.2)
        bp.paragraph_format.space_after = Pt(3)
        rb = bp.add_run(f"✔  {title_b}")
        rb.bold = True
        rb.font.color.rgb = COLOR_PRIMARY
        bp.add_run(desc_b)

    doc.save(filename)
    print(f"Documento guardado con éxito en: {filename}")

if __name__ == '__main__':
    target = os.path.join(os.getcwd(), 'Capacitacion_MongoDB_Mongoose.docx')
    create_training_document(target)
